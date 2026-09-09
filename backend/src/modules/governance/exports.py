import csv
import hashlib
import io
import uuid
from datetime import UTC
from decimal import Decimal
from typing import Any

from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from ...infrastructure.config.settings import get_settings
from .exceptions import ReportNotFoundError, ReportValidationError
from .models import Artifact, AuditEvent, ExportJob, OutboxEvent, ReportVersion, SnapshotItem, now
from .service import canonical_json
from .storage import ReportArtifactStorage, safe_filename

MIME_TYPES = {
    "CSV": "text/csv; charset=utf-8",
    "PDF": "application/pdf",
    "PPTX": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def _emit(
    db: AsyncSession, job: ExportJob, event_type: str, actor_id: int | None, detail: dict[str, Any] | None = None
) -> None:
    body = {
        "report_version_id": str(job.report_version_id),
        "export_job_id": str(job.id),
        "actor_id": actor_id,
        "format": job.format,
        "occurred_at": now().isoformat(),
        **(detail or {}),
    }
    db.add(
        AuditEvent(
            event_type=event_type,
            entity_type="EXPORT_JOB",
            entity_id=job.id,
            payload=body,
            correlation_id=str(job.id),
            actor_id=actor_id,
        )
    )
    db.add(OutboxEvent(event_type=event_type, aggregate_id=job.id, payload=body, available_at=now()))


async def request_export(
    db: AsyncSession,
    *,
    report_version_id: uuid.UUID,
    format_: str,
    options: dict[str, Any],
    template_version: str,
    actor_id: int,
    retry_failed: bool,
) -> tuple[ExportJob, bool]:
    version = await db.get(ReportVersion, report_version_id)
    if version is None or version.published_at is None:
        raise ReportNotFoundError("Published report version not found")
    if len(canonical_json(options).encode("utf-8")) > 8192:
        raise ReportValidationError("Export options are too large")
    identity = {
        "report_version_id": str(report_version_id),
        "format": format_,
        "options": options,
        "template_version": template_version,
    }
    key = hashlib.sha256(canonical_json(identity).encode("utf-8")).hexdigest()
    existing = (await db.scalars(select(ExportJob).where(ExportJob.idempotency_key == key))).one_or_none()
    if existing:
        if existing.status == "FAILED" and retry_failed:
            existing.status = "QUEUED"
            existing.error_message = None
            existing.finished_at = None
            existing.updated_at = now()
            _emit(db, existing, "REPORT_EXPORT_REQUESTED", actor_id, {"retry": True})
            await db.commit()
            return existing, True
        return existing, False
    job = ExportJob(
        report_version_id=report_version_id,
        idempotency_key=key,
        format=format_,
        status="QUEUED",
        options=options,
        template_version=template_version,
        requested_by_user_id=actor_id,
    )
    db.add(job)
    try:
        await db.flush()
    except IntegrityError:
        await db.rollback()
        raced = (await db.scalars(select(ExportJob).where(ExportJob.idempotency_key == key))).one()
        return raced, False
    _emit(db, job, "REPORT_EXPORT_REQUESTED", actor_id)
    await db.commit()
    await db.refresh(job)
    return job, True


def job_dict(job: ExportJob, artifact: Artifact | None = None) -> dict[str, Any]:
    return {
        "id": str(job.id),
        "report_version_id": str(job.report_version_id),
        "format": job.format,
        "status": job.status,
        "attempts": job.attempts,
        "created_at": job.created_at.isoformat(),
        "updated_at": job.updated_at.isoformat(),
        "started_at": job.started_at.isoformat() if job.started_at else None,
        "finished_at": job.finished_at.isoformat() if job.finished_at else None,
        "error_message": job.error_message,
        "artifact": (
            {
                "filename": artifact.filename,
                "content_type": artifact.content_type,
                "size_bytes": artifact.size_bytes,
                "sha256": artifact.sha256,
                "download_url": f"/api/v1/exports/{job.id}/download",
            }
            if artifact
            else None
        ),
    }


def _csv_text(value: Any) -> str:
    text = "" if value is None else str(value)
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def render_csv(version: ReportVersion, items: list[dict[str, Any]]) -> bytes:
    buffer = io.StringIO(newline="")
    fields = [
        "occurrence_id",
        "transaction_date",
        "organization_code",
        "item_code",
        "item_description",
        "product",
        "division",
        "line",
        "issue_quantity",
        "issue_amount_brl",
        "amount_usd",
        "exchange_rate",
        "exchange_rate_effective_date",
        "review_title",
        "review_description",
        "reviewed_by_name",
        "reviewed_at",
    ]
    writer = csv.DictWriter(buffer, fieldnames=fields, extrasaction="ignore", lineterminator="\r\n")
    writer.writeheader()
    numeric = {"issue_quantity", "issue_amount_brl", "amount_usd", "exchange_rate"}
    for item in items:
        writer.writerow(
            {
                field: (
                    str(Decimal(str(item[field])))
                    if field in numeric and item.get(field) is not None
                    else _csv_text(item.get(field))
                )
                for field in fields
            }
        )
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def render_pdf(version: ReportVersion, items: list[dict[str, Any]]) -> bytes:
    from reportlab.lib import colors  # noqa: PLC0415
    from reportlab.lib.pagesizes import A4, landscape  # noqa: PLC0415
    from reportlab.lib.styles import getSampleStyleSheet  # noqa: PLC0415
    from reportlab.lib.units import mm  # noqa: PLC0415
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: PLC0415

    output = io.BytesIO()
    report = version.content.get("report", {})
    styles = getSampleStyleSheet()

    def footer(canvas: Any, document: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(15 * mm, 9 * mm, f"{report.get('code', 'REPORT')} · rev. {version.revision} · {version.sha256[:16]}")
        canvas.drawRightString(282 * mm, 9 * mm, f"Page {document.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=14 * mm,
        bottomMargin=15 * mm,
        title=str(report.get("title", "Scrap report")),
    )
    story = [
        Paragraph(str(report.get("title", "Scrap report")), styles["Title"]),
        Paragraph(f"{report.get('code', '')} · Revision {version.revision}", styles["Heading2"]),
        Spacer(1, 6 * mm),
        Paragraph(str(report.get("description", "")), styles["BodyText"]),
        Spacer(1, 5 * mm),
    ]
    metrics = version.content.get("metrics", {})
    summary_text = (
        f"Occurrences: {metrics.get('occurrence_count', 0)} · "
        f"BRL {metrics.get('issue_amount_brl', '0')} · USD {metrics.get('amount_usd', '0')}"
    )
    story.extend(
        [
            Paragraph("Summary", styles["Heading2"]),
            Paragraph(summary_text, styles["BodyText"]),
            PageBreak(),
        ]
    )
    data = [["Date", "Organization", "Item", "Description", "USD", "Review", "Author"]]
    for item in items:
        data.append(
            [
                item.get("transaction_date", ""),
                item.get("organization_code", ""),
                item.get("item_code", ""),
                Paragraph(str(item.get("item_description") or ""), styles["BodyText"]),
                item.get("amount_usd", ""),
                Paragraph(str(item.get("review_title") or ""), styles["BodyText"]),
                item.get("reviewed_by_name", ""),
            ]
        )
    table = Table(data, repeatRows=1, colWidths=[25 * mm, 28 * mm, 30 * mm, 62 * mm, 27 * mm, 66 * mm, 38 * mm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12372A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D8E2DC")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F8F6")]),
            ]
        )
    )
    story.append(table)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def render_pptx(version: ReportVersion, items: list[dict[str, Any]]) -> bytes:
    from pptx import Presentation  # noqa: PLC0415
    from pptx.dml.color import RGBColor  # noqa: PLC0415
    from pptx.util import Inches, Pt  # noqa: PLC0415

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    report = version.content.get("report", {})
    metrics = version.content.get("metrics", {})
    cover = presentation.slides.add_slide(presentation.slide_layouts[0])
    cover.shapes.title.text = str(report.get("title", "Scrap report"))
    published = version.published_at.astimezone(UTC)
    cover.placeholders[1].text = (
        f"{report.get('code', '')} · Revision {version.revision}\n"
        f"Published {published:%Y-%m-%d %H:%M UTC}\nTrace {version.sha256[:16]}"
    )
    summary = presentation.slides.add_slide(presentation.slide_layouts[5])
    summary.shapes.title.text = "Executive summary"
    box = summary.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(2.3))
    frame = box.text_frame
    frame.text = f"{metrics.get('occurrence_count', 0)} occurrences"
    for text in (f"BRL {metrics.get('issue_amount_brl', '0')}", f"USD {metrics.get('amount_usd', '0')}"):
        paragraph = frame.add_paragraph()
        paragraph.text = text
        paragraph.font.size = Pt(28)
        paragraph.font.color.rgb = RGBColor(18, 55, 42)
    for offset in range(0, len(items), 6):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = f"Occurrences {offset + 1}–{min(offset + 6, len(items))}"
        rows = min(6, len(items) - offset) + 1
        shape = slide.shapes.add_table(rows, 5, Inches(0.45), Inches(1.35), Inches(12.4), Inches(5.4))
        table = shape.table
        for column, heading in enumerate(("Date", "Organization", "Item", "USD", "Review")):
            table.cell(0, column).text = heading
        for row, item in enumerate(items[offset : offset + 6], start=1):
            for column, value in enumerate(
                (
                    item.get("transaction_date"),
                    item.get("organization_code"),
                    item.get("item_code"),
                    item.get("amount_usd"),
                    item.get("review_title"),
                )
            ):
                table.cell(row, column).text = str(value or "")
        notes = slide.notes_slide.notes_text_frame
        notes.text = f"Source: immutable Hanaro snapshot · report revision {version.revision} · {version.sha256}"
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


async def run_export(db: AsyncSession, job_id: uuid.UUID, storage: ReportArtifactStorage | None = None) -> ExportJob:
    job = await db.get(ExportJob, job_id, with_for_update=True)
    if job is None:
        raise ReportNotFoundError("Export job not found")
    existing = (await db.scalars(select(Artifact).where(Artifact.export_job_id == job.id))).one_or_none()
    storage = storage or ReportArtifactStorage()
    if existing:
        try:
            content = storage.read(existing.storage_key)
        except (FileNotFoundError, ValueError):
            pass
        else:
            if hashlib.sha256(content).hexdigest() == existing.sha256:
                job.status = "COMPLETED"
                job.finished_at = job.updated_at = now()
                await db.commit()
                return job
    job.status = "RUNNING"
    job.attempts += 1
    job.started_at = job.updated_at = now()
    job.error_message = None
    await db.commit()
    try:
        version = await db.get(ReportVersion, job.report_version_id)
        if version is None or version.published_at is None:
            raise ReportNotFoundError("Published report version not found")
        items = [
            item.frozen_values
            for item in await db.scalars(
                select(SnapshotItem).where(SnapshotItem.snapshot_id == version.snapshot_id).order_by(SnapshotItem.occurrence_id)
            )
        ]
        if len(items) > get_settings().REPORT_EXPORT_MAX_ITEMS:
            raise ReportValidationError("Report exceeds the configured export item limit")
        renderers = {"CSV": render_csv, "PDF": render_pdf, "PPTX": render_pptx}
        content = renderers[job.format](version, items)
        digest = hashlib.sha256(content).hexdigest()
        extension = job.format.lower()
        filename = safe_filename(f"report-{version.report_id}-rev-{version.revision}.{extension}")
        key = f"{version.report_id}/{version.id}/{job.id}.{extension}"
        storage.write_once(key, content)
        artifact = Artifact(
            export_job_id=job.id,
            storage_key=key,
            sha256=digest,
            size_bytes=len(content),
            filename=filename,
            content_type=MIME_TYPES[job.format],
        )
        db.add(artifact)
        job.status = "COMPLETED"
        job.finished_at = job.updated_at = now()
        _emit(db, job, "REPORT_EXPORT_COMPLETED", job.requested_by_user_id, {"sha256": digest, "size_bytes": len(content)})
    except Exception as error:
        job.status = "FAILED"
        job.finished_at = job.updated_at = now()
        job.error_message = "The report export could not be generated"
        _emit(db, job, "REPORT_EXPORT_FAILED", job.requested_by_user_id, {"error_type": type(error).__name__})
        await db.commit()
        raise
    await db.commit()
    return job
