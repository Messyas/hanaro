import csv
import importlib.util
import io
import uuid
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from unittest.mock import AsyncMock

import pytest
from alembic.migration import MigrationContext
from alembic.operations import Operations
from PIL import Image
from pptx import Presentation
from pydantic import ValidationError
from sqlalchemy import create_engine, func, inspect, select, text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from src.infrastructure.database.session import Base
from src.modules.governance.action_evidence import add_evidence, read_evidence, remove_evidence, validate_file
from src.modules.governance.actions import PlanInput, TaskCommand, TaskInput, board, command_task, save_plan, save_task
from src.modules.governance.exceptions import ReportConflictError, ReportNotFoundError, ReportValidationError
from src.modules.governance.exports import request_export, run_export
from src.modules.governance.exports.document import Document
from src.modules.governance.exports.renderers import RENDERERS
from src.modules.governance.models import (
    ActionEvidence,
    ActionPlan,
    Alert,
    AlertRecipient,
    Artifact,
    AuditEvent,
    ConsumerReceipt,
    EmailDelivery,
    EventAttempt,
    Factory,
    OutboxEvent,
    PlanReport,
    ReportVersion,
)
from src.modules.governance.notifications.monitor import monitor
from src.modules.governance.notifications.rules import RuleInput, evaluate, save_rule
from src.modules.governance.notifications.service import consume, dispatch, emit, recover_exports
from src.modules.governance.schemas import ExportOptions
from src.modules.governance.service import create_report, mutate_sources, publish_report
from src.modules.governance.storage import ReportArtifactStorage
from src.modules.governance.workflow_routes import history as list_history
from src.modules.governance.workflow_routes import plans as list_plans
from src.modules.material_scrap.models import ScrapOccurrence, ScrapReview
from src.modules.material_scrap.service import ingest_material_scrap
from src.modules.user.models import User

from ..material_scrap.helpers import canonical_fixture


@pytest.fixture
async def db():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    async with async_sessionmaker(engine, expire_on_commit=False)() as session:
        session.add_all(
            [
                User(name="Analyst", username="analyst", email="one@example.com", hashed_password="hash"),
                User(name="Peer", username="peer", email="two@example.com", hashed_password="hash"),
            ]
        )
        await session.commit()
        yield session
    await engine.dispose()


async def published(db):
    await ingest_material_scrap(canonical_fixture(), db)
    occurrence = await db.scalar(select(ScrapOccurrence))
    review = ScrapReview(
        occurrence_id=occurrence.id,
        responsible_user_id=1,
        responsible_name="Analyst",
        status="REVIEWED",
        title="Review <&>",
        description="Complete analysis\n" * 150,
        version=1,
        created_at=datetime.now(UTC),
        updated_at=datetime.now(UTC),
    )
    db.add(review)
    await db.commit()
    report = await create_report(
        db, title="Relatório <&> 한글", description="Long report", factory_id=None, actor_id=1, correlation_id="test"
    )
    await mutate_sources(
        db,
        report.id,
        kind="occurrence",
        operation="add",
        ids=[occurrence.id],
        expected_version=report.version,
        actor_id=1,
        correlation_id="test",
    )
    return await publish_report(
        db, report.id, expected_version=report.version, template_version="1", actor_id=1, correlation_id="test"
    )


async def test_task_validation_participants_conflict_and_history(db):
    version = await published(db)
    plan = await save_plan(db, PlanInput(title="Corrective actions", report_version_ids=[version.id]), 1)
    task = await save_task(db, plan["id"], TaskInput(title="Repair", participant_ids=[1, 2], blocked_reason="Waiting"), 1)
    assert task["is_blocked"] is True
    assert {p["id"] for p in task["participants"]} == {1, 2}
    assert (await board(db, plan["id"], "PLANNED", 1, 1))["total"] == 1
    with pytest.raises(ReportConflictError):
        await command_task(db, task["id"], TaskCommand(expected_version=task["version"], command="validate"), 2)
    task = await command_task(
        db, task["id"], TaskCommand(expected_version=task["version"], command="move", status="UNDER_VERIFICATION"), 2
    )
    with pytest.raises(ReportConflictError):
        await command_task(db, task["id"], TaskCommand(expected_version=task["version"], command="validate"), 1)
    task = await save_task(
        db, plan["id"], TaskInput(title="Repair", participant_ids=[1, 2], expected_version=task["version"]), 2, task["id"]
    )
    assert task["is_blocked"] is False
    task = await command_task(db, task["id"], TaskCommand(expected_version=task["version"], command="validate"), 2)
    assert task["status"] == "COMPLETED" and task["validated_by_id"] == 2
    with pytest.raises(ReportConflictError):
        await command_task(db, task["id"], TaskCommand(expected_version=1, command="reopen"), 1)
    task = await command_task(db, task["id"], TaskCommand(expected_version=task["version"], command="reopen"), 1)
    assert task["status"] == "IN_PROGRESS"


async def test_board_filters_participants_and_tags_without_duplicate_tasks(db):
    version = await published(db)
    plan = await save_plan(db, PlanInput(title="Corrective actions", report_version_ids=[version.id]), 1)
    matching = await save_task(
        db,
        plan["id"],
        TaskInput(title="Repair", participant_ids=[1, 2], tags=["Supplier", "Urgent"]),
        1,
    )
    await save_task(db, plan["id"], TaskInput(title="Inspect", participant_ids=[1], tags=["Other"]), 1)

    page = await board(db, plan["id"], "PLANNED", 1, 1, participant="peer", tag="supp")
    assert page["total"] == 1
    assert page["items"][0]["id"] == matching["id"]
    assert page["has_next"] is False
    assert (await board(db, plan["id"], "PLANNED", 1, 25, participant="analyst"))["total"] == 2
    assert (await board(db, plan["id"], "PLANNED", 1, 25, tag="suppliers"))["total"] == 0


async def test_block_without_reason_keeps_kanban_column_and_emits_once(db):
    version = await published(db)
    plan = await save_plan(db, PlanInput(title="Corrective actions", report_version_ids=[version.id]), 1)
    task = await save_task(db, plan["id"], TaskInput(title="Repair", participant_ids=[2], is_blocked=True), 1)
    assert task["is_blocked"] is True and task["blocked_reason"] is None and task["status"] == "PLANNED"
    event = await db.scalar(select(OutboxEvent).where(OutboxEvent.event_type == "TASK_BLOCKED"))
    assert event.payload["recipient_ids"] == [2]
    await consume(db, event)
    await db.commit()
    assert await db.scalar(select(Alert).where(Alert.event_id == event.id)) is not None
    assert await db.scalar(select(EmailDelivery).where(EmailDelivery.event_id == event.id)) is not None
    audit = await db.scalar(select(AuditEvent).where(AuditEvent.entity_id == task["id"], AuditEvent.event_type == "TASK_SAVED"))
    assert audit.payload["is_blocked"] is True

    task = await save_task(
        db,
        plan["id"],
        TaskInput(title="Repair", participant_ids=[2], is_blocked=True, expected_version=task["version"]),
        1,
        task["id"],
    )
    assert await db.scalar(select(func.count()).select_from(OutboxEvent).where(OutboxEvent.event_type == "TASK_BLOCKED")) == 1
    task = await command_task(
        db, task["id"], TaskCommand(expected_version=task["version"], command="move", status="UNDER_VERIFICATION"), 1
    )
    with pytest.raises(ReportConflictError):
        await command_task(db, task["id"], TaskCommand(expected_version=task["version"], command="validate"), 2)

    task = await save_task(
        db, plan["id"], TaskInput(title="Repair", participant_ids=[2], expected_version=task["version"]), 1, task["id"]
    )
    assert task["is_blocked"] is False and task["status"] == "UNDER_VERIFICATION"
    task = await command_task(db, task["id"], TaskCommand(expected_version=task["version"], command="validate"), 2)
    assert task["status"] == "COMPLETED"
    with pytest.raises(ValidationError):
        TaskInput(title="Invalid", is_blocked=False, blocked_reason="Waiting")


async def test_task_deadline_events_fire_without_rules_once_per_deadline_and_stage(db):
    version = await published(db)
    plan = await save_plan(db, PlanInput(title="Corrective actions", report_version_ids=[version.id]), 1)
    due = datetime.now(UTC) + timedelta(hours=2)
    task = await save_task(db, plan["id"], TaskInput(title="Repair", due_at=due, participant_ids=[2]), 1)

    await monitor(db)
    await monitor(db)
    due_events = list(
        await db.scalars(
            select(OutboxEvent).where(OutboxEvent.aggregate_id == task["id"], OutboxEvent.event_type == "TASK_DUE")
        )
    )
    assert len(due_events) == 1
    assert due_events[0].payload["recipient_ids"] == [2]
    assert due_events[0].payload["due_at"] == due.isoformat()
    await consume(db, due_events[0])
    await db.commit()
    assert (
        await db.scalar(select(func.count()).select_from(EmailDelivery).where(EmailDelivery.event_id == due_events[0].id)) == 1
    )

    past_due = datetime.now(UTC) - timedelta(hours=1)
    task = await save_task(
        db,
        plan["id"],
        TaskInput(title="Repair", due_at=past_due, participant_ids=[2], expected_version=task["version"]),
        1,
        task["id"],
    )
    await monitor(db)
    await monitor(db)
    assert (
        await db.scalar(
            select(func.count())
            .select_from(OutboxEvent)
            .where(OutboxEvent.aggregate_id == task["id"], OutboxEvent.event_type == "TASK_OVERDUE")
        )
        == 1
    )

    untouched = await save_task(db, plan["id"], TaskInput(title="Inspect", due_at=due, participant_ids=[2]), 1)
    await save_plan(
        db,
        PlanInput(
            title="Corrective actions",
            report_version_ids=[version.id],
            expected_version=plan["version"],
            status="COMPLETED",
        ),
        2,
        plan["id"],
    )
    await save_rule(db, RuleInput(name="Overdue reminder", event_type="TASK_OVERDUE", user_ids=[2]))
    await monitor(db)
    assert (
        await db.scalar(select(func.count()).select_from(OutboxEvent).where(OutboxEvent.aggregate_id == untouched["id"])) == 1
    )  # Assignment only; completed plans do not generate deadline events.
    assert (
        await db.scalar(
            select(func.count())
            .select_from(OutboxEvent)
            .where(OutboxEvent.aggregate_id == task["id"], OutboxEvent.event_type == "TASK_OVERDUE")
        )
        == 1
    )


async def test_configured_due_rule_does_not_duplicate_default_deadline_event(db):
    version = await published(db)
    plan = await save_plan(db, PlanInput(title="Corrective actions", report_version_ids=[version.id]), 1)
    task = await save_task(
        db,
        plan["id"],
        TaskInput(title="Repair", due_at=datetime.now(UTC) + timedelta(hours=2), participant_ids=[2]),
        1,
    )
    rule = await save_rule(db, RuleInput(name="Due reminder", event_type="TASK_DUE", user_ids=[2]))

    await monitor(db)
    await monitor(db)
    events = list(
        await db.scalars(
            select(OutboxEvent).where(OutboxEvent.aggregate_id == task["id"], OutboxEvent.event_type == "TASK_DUE")
        )
    )
    assert len(events) == 1
    assert events[0].payload["rule_id"] == str(rule["id"])


async def test_task_occurrences_and_comments_are_available_in_detail_and_history(db):
    version = await published(db)
    occurrence = await db.scalar(select(ScrapOccurrence))
    plan = await save_plan(db, PlanInput(title="Corrective actions", report_version_ids=[version.id]), 1)
    task = await save_task(
        db,
        plan["id"],
        TaskInput(title="Repair", occurrence_ids=[occurrence.id, occurrence.id]),
        1,
    )
    assert task["occurrence_ids"] == [occurrence.id]
    assert task["occurrences"][0]["id"] == occurrence.id
    assert task["occurrences"][0]["item_code"]

    task = await command_task(
        db, task["id"], TaskCommand(expected_version=task["version"], command="comment", comment="Check supplier"), 2
    )
    events = await list_history(task["id"], db, {"id": 1})
    assert events["items"][0]["payload"]["comment"] == "Check supplier"
    assert events["items"][0]["actor_name"] == "Peer"

    task = await command_task(
        db, task["id"], TaskCommand(expected_version=task["version"], command="move", status="UNDER_VERIFICATION"), 1
    )
    task = await command_task(
        db, task["id"], TaskCommand(expected_version=task["version"], command="validate", comment="Verified"), 2
    )
    assert task["status"] == "COMPLETED"
    events = await list_history(task["id"], db, {"id": 1})
    assert events["items"][0]["payload"]["comment"] == "Verified"
    task = await command_task(db, task["id"], TaskCommand(expected_version=task["version"], command="reopen"), 1)

    task = await save_task(db, plan["id"], TaskInput(title="Repair", expected_version=task["version"]), 1, task["id"])
    assert task["occurrences"] == [] and task["occurrence_ids"] == []


async def test_task_tags_and_evidence_are_versioned_private_and_audited(db, tmp_path):
    version = await published(db)
    plan = await save_plan(db, PlanInput(title="Corrective actions", report_version_ids=[version.id]), 1)
    task = await save_task(db, plan["id"], TaskInput(title="Repair", tags=[" Urgent ", "urgent", "Supplier"]), 1)
    assert task["tags"] == ["Urgent", "Supplier"]
    assert task["evidence"] == []

    content = io.BytesIO()
    Image.new("RGB", (2, 2), "red").save(content, format="PNG")
    storage = ReportArtifactStorage(tmp_path)
    task = await add_evidence(
        db,
        task["id"],
        content=content.getvalue(),
        filename="proof.html",
        content_type="image/png",
        expected_version=task["version"],
        actor_id=2,
        storage=storage,
    )
    evidence_id = task["evidence"][0]["id"]
    assert task["evidence"][0]["filename"] == "proof.png"
    assert task["version"] > 1
    stored, downloaded = await read_evidence(db, task["id"], evidence_id, storage)
    assert stored.uploaded_by_user_id == 2 and downloaded == content.getvalue()
    with pytest.raises(ReportConflictError):
        await remove_evidence(db, task["id"], evidence_id, 1, 1)
    await db.rollback()

    task = await remove_evidence(db, task["id"], evidence_id, task["version"], 1)
    assert task["evidence"] == []
    assert await db.get(ActionEvidence, evidence_id) is not None
    with pytest.raises(ReportNotFoundError):
        await read_evidence(db, task["id"], evidence_id, storage)
    audits = list(
        await db.scalars(
            select(AuditEvent)
            .where(AuditEvent.entity_id == task["id"], AuditEvent.event_type.like("TASK_EVIDENCE_%"))
            .order_by(AuditEvent.created_at)
        )
    )
    assert {event.event_type for event in audits} == {"TASK_EVIDENCE_ADDED", "TASK_EVIDENCE_REMOVED"}


def test_task_evidence_rejects_unsafe_content_and_tags():
    with pytest.raises(ReportValidationError):
        validate_file(b"not-an-image", "proof.png", "image/png")
    with pytest.raises(ReportValidationError):
        validate_file(b"%PDF-1.7", "proof.pdf", "application/pdf")
    with pytest.raises(ReportValidationError):
        validate_file(b"x" * 10_000_001, "proof.pdf", "application/pdf")
    with pytest.raises(ValidationError):
        TaskInput(title="Invalid", tags=["x" * 41])


async def test_action_plan_list_includes_empty_reports_contract(db):
    factory = Factory(code="TEST", name="Test factory")
    db.add(factory)
    await db.flush()
    created = ActionPlan(factory_id=factory.id, title="Legacy plan", author_id=1)
    db.add(created)
    await db.commit()

    page = await list_plans(db, {"id": 1}, page=1, page_size=25)

    assert page["items"][0]["id"] == created.id
    assert page["items"][0]["reports"] == []


async def test_action_plan_list_filters_and_orders_pages_deterministically(db):
    factory = Factory(code="FILTER", name="Filter factory")
    db.add(factory)
    await db.flush()
    created_at = datetime(2026, 9, 1, tzinfo=UTC)
    plans = [
        ActionPlan(id=uuid.UUID(int=1), factory_id=factory.id, title="Alpha", status="OPEN", created_at=created_at),
        ActionPlan(id=uuid.UUID(int=2), factory_id=factory.id, title="Alpha", status="OPEN", created_at=created_at),
        ActionPlan(
            id=uuid.UUID(int=3),
            factory_id=factory.id,
            title="Alpha",
            status="COMPLETED",
            created_at=created_at + timedelta(days=1),
        ),
        ActionPlan(
            id=uuid.UUID(int=4),
            factory_id=factory.id,
            title="Beta",
            status="OPEN",
            created_at=created_at + timedelta(days=2),
        ),
    ]
    db.add_all(plans)
    await db.commit()

    first = await list_plans(db, {"id": 1}, page=1, page_size=1, search=" alpha ", status="OPEN")
    second = await list_plans(db, {"id": 1}, page=2, page_size=1, search="alpha", status="OPEN")
    assert first["total"] == second["total"] == 2
    assert first["has_next"] is True and second["has_next"] is False
    assert [first["items"][0]["id"], second["items"][0]["id"]] == [plans[0].id, plans[1].id]

    oldest = await list_plans(db, {"id": 1}, page=1, page_size=10, sort="oldest")
    newest = await list_plans(db, {"id": 1}, page=1, page_size=10, sort="newest")
    title = await list_plans(db, {"id": 1}, page=1, page_size=10, sort="title")
    assert [item["id"] for item in oldest["items"]] == [plan.id for plan in plans]
    assert [item["id"] for item in newest["items"]] == [plans[3].id, plans[2].id, plans[0].id, plans[1].id]
    assert [item["id"] for item in title["items"]] == [plans[0].id, plans[1].id, plans[2].id, plans[3].id]


async def test_action_plan_requires_a_report_and_deduplicates_links(db):
    with pytest.raises(ValidationError):
        PlanInput(title="No source", report_version_ids=[])

    first = await published(db)
    second_report = await create_report(
        db, title="Second report", description="", factory_id=None, actor_id=1, correlation_id="test"
    )
    occurrence_id = await db.scalar(select(ScrapReview.occurrence_id))
    await mutate_sources(
        db,
        second_report.id,
        kind="occurrence",
        operation="add",
        ids=[occurrence_id],
        expected_version=second_report.version,
        actor_id=1,
        correlation_id="test",
    )
    second = await publish_report(
        db, second_report.id, expected_version=second_report.version, template_version="1", actor_id=1, correlation_id="test"
    )
    plan = await save_plan(
        db,
        PlanInput(title="Two reports", report_version_ids=[first.id, first.id, second.id]),
        1,
    )
    assert {link["report_id"] for link in plan["reports"]} == {first.report_id, second.report_id}
    assert await db.scalar(select(func.count()).select_from(PlanReport)) == 2


async def test_action_plan_rejects_two_versions_of_one_report(db):
    first = await published(db)
    second = ReportVersion(
        report_id=first.report_id,
        snapshot_id=first.snapshot_id,
        revision=first.revision + 1,
        content=first.content,
        template_version=first.template_version,
        content_schema_version=first.content_schema_version,
        sha256=first.sha256,
        published_at=datetime.now(UTC),
    )
    db.add(second)
    await db.commit()

    with pytest.raises(ReportValidationError, match="one version"):
        await save_plan(db, PlanInput(title="Duplicate report", report_version_ids=[first.id, second.id]), 1)
    await db.rollback()
    assert await db.scalar(select(func.count()).select_from(ActionPlan)) == 0


async def test_email_simulation_is_idempotent_and_separate_from_alerts(db):
    event = emit(db, "TASK_ASSIGNED", uuid.uuid4(), {"title": "Repair", "recipient_ids": [1, 2]})
    await db.flush()
    await consume(db, event)
    await db.commit()
    await consume(db, event)
    await db.commit()
    assert await db.scalar(select(func.count()).select_from(EmailDelivery)) == 2
    assert await db.scalar(select(func.count()).select_from(Alert)) == 0
    assert await db.scalar(select(func.count()).select_from(ConsumerReceipt)) == 1
    assert (await db.scalar(select(EmailDelivery))).status == "SIMULATED"
    event = emit(db, "TASK_OVERDUE", uuid.uuid4(), {"title": "Late", "recipient_ids": [1, 2]})
    await db.flush()
    await consume(db, event)
    await db.commit()
    recipients = list(await db.scalars(select(AlertRecipient).order_by(AlertRecipient.user_id)))
    recipients[0].read_at = datetime.now(UTC)
    await db.commit()
    assert recipients[1].read_at is None
    assert (await db.scalar(select(Alert))).resolved_at is None


async def test_broker_failure_is_recorded_and_recoverable(db):
    event = emit(db, "REPORT_EXPORT_REQUESTED", uuid.uuid4(), {})
    await db.commit()
    broker = AsyncMock(side_effect=RuntimeError("secret connection string"))
    assert await dispatch(db, broker) == 0
    await db.refresh(event)
    assert event.published_at is None and event.attempts == 1
    attempt = await db.scalar(select(EventAttempt))
    assert "secret" not in attempt.error
    event.available_at = datetime.now(UTC) - timedelta(seconds=1)
    await db.commit()
    assert await dispatch(db, AsyncMock()) == 1


async def test_exports_share_identity_and_expired_claim_is_recovered(db, tmp_path):
    version = await published(db)
    args = dict(report_version_id=version.id, format_="MARKDOWN", options={}, template_version="1", retry_failed=False)
    first, _ = await request_export(db, actor_id=1, **args)
    second, created = await request_export(db, actor_id=2, **{**args, "options": ExportOptions().model_dump()})
    assert first.id == second.id and not created
    storage = ReportArtifactStorage(tmp_path)
    result = await run_export(db, first.id, storage)
    assert result.status == "COMPLETED"
    artifact = await db.scalar(select(Artifact))
    original = storage.read(artifact.storage_key)
    await run_export(db, first.id, storage)
    assert storage.read(artifact.storage_key) == original
    assert result.attempts == 1
    completed = await db.scalar(select(OutboxEvent).where(OutboxEvent.event_type == "REPORT_EXPORT_COMPLETED"))
    await consume(db, completed)
    await db.commit()
    assert await db.scalar(select(func.count()).select_from(EmailDelivery)) == 0
    notified, _ = await request_export(
        db,
        actor_id=1,
        **{
            **args,
            "format_": "CSV",
            "options": ExportOptions(notify_on_completion=True).model_dump(),
        },
    )
    await run_export(db, notified.id, storage)
    notified_event = await db.scalar(
        select(OutboxEvent).where(
            OutboxEvent.event_type == "REPORT_EXPORT_COMPLETED",
            OutboxEvent.aggregate_id == notified.id,
        )
    )
    await consume(db, notified_event)
    await db.commit()
    assert await db.scalar(select(func.count()).select_from(EmailDelivery)) == 1
    other, _ = await request_export(db, actor_id=1, **{**args, "format_": "PDF"})
    other.status = "RUNNING"
    other.lease_until = datetime.now(UTC) - timedelta(seconds=1)
    other.lease_token = uuid.uuid4()
    await db.commit()
    await recover_exports(db)
    assert other.status == "QUEUED" and other.lease_token is None


async def test_rules_dedupe_and_closed_cost_goal(db):
    await published(db)
    await save_rule(db, RuleInput(name="Cost", event_type="SCRAP_RELEVANT", threshold=Decimal(0), user_ids=[1]))
    await evaluate(db, today=date(2026, 9, 9))
    before = await db.scalar(select(func.count()).select_from(OutboxEvent).where(OutboxEvent.event_type == "SCRAP_RELEVANT"))
    await evaluate(db, today=date(2026, 9, 9))
    assert (
        await db.scalar(select(func.count()).select_from(OutboxEvent).where(OutboxEvent.event_type == "SCRAP_RELEVANT"))
        == before
    )
    with pytest.raises(ValidationError):
        RuleInput(name="Bad goal", event_type="COST_EXCEEDED", severity="POSITIVE")
    with pytest.raises(ValidationError):
        ExportOptions(unsafe=True)


def test_real_formats_long_text_and_options(tmp_path):
    version = ReportVersion(
        report_id=uuid.uuid4(),
        snapshot_id=uuid.uuid4(),
        revision=1,
        template_version="1",
        sha256="f" * 64,
        published_at=datetime.now(UTC),
        content={
            "report": {"title": "Análise <&> 한글", "author": "Analyst"},
            "metrics": {"occurrence_count": 1, "amount_usd": "-12.34", "issue_amount_brl": "-50.00"},
        },
    )
    item = {
        "occurrence_id": str(uuid.uuid4()),
        "item_code": " \t=SUM(1,2)",
        "issue_quantity": "-12.340000",
        "amount_usd": "-12.34",
        "review_description": "Justificativa completa 한글 <script>\n" * 120,
        "review_title": "Title",
    }
    for language in ("pt", "en", "ko"):
        for name, renderer in RENDERERS.items():
            content = renderer(Document(version, [item], ExportOptions(language=language)))
            (tmp_path / f"report-{language}.{name.lower()}").write_bytes(content)
            if name == "PPTX":
                slides = Presentation(io.BytesIO(content)).slides
                assert len(slides) > 8
                assert "Justificativa completa" in " ".join(
                    shape.text for slide in slides for shape in slide.shapes if shape.has_text_frame
                )
            if name == "CSV":
                row = next(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
                assert row["item_code"].startswith("'") and row["issue_quantity"] == "-12.340000"
            if name == "MARKDOWN":
                assert b"<script>" not in content
    minimal = RENDERERS["MARKDOWN"](Document(version, [item], ExportOptions(include_money=False, include_justifications=False)))
    assert b"12.34" not in minimal and b"Justificativa completa" not in minimal


async def test_migration_reconciles_precreated_schema(db):
    path = Path(__file__).parents[4] / "migrations/versions/20260909_14_governance_workflows.py"
    spec = importlib.util.spec_from_file_location("workflow_migration", path)
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    connection = await db.connection()

    def upgrade(connection):
        with Operations.context(MigrationContext.configure(connection)):
            migration.upgrade()

    await connection.run_sync(upgrade)
    await db.commit()


def test_block_migration_backfills_legacy_reason_and_preserves_rollback():
    path = Path(__file__).parents[4] / "migrations/versions/20260924_25_action_blocking.py"
    spec = importlib.util.spec_from_file_location("block_migration", path)
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    engine = create_engine("sqlite:///:memory:")
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE gov_actions (id INTEGER PRIMARY KEY, blocked_reason TEXT)"))
        connection.execute(text("INSERT INTO gov_actions (id, blocked_reason) VALUES (1, 'Waiting'), (2, NULL)"))
        with Operations.context(MigrationContext.configure(connection)):
            migration.upgrade()
        assert connection.execute(text("SELECT id, is_blocked FROM gov_actions ORDER BY id")).all() == [(1, 1), (2, 0)]
        connection.execute(text("UPDATE gov_actions SET is_blocked = true WHERE id = 2"))
        with Operations.context(MigrationContext.configure(connection)):
            migration.downgrade()
        assert connection.execute(text("SELECT blocked_reason FROM gov_actions WHERE id = 2")).scalar() == (
            "Bloqueada sem motivo informado."
        )
    engine.dispose()


def test_task_evidence_migration_preserves_existing_data_on_rollback():
    path = Path(__file__).parents[4] / "migrations/versions/20260924_26_action_tags_evidence.py"
    spec = importlib.util.spec_from_file_location("task_evidence_migration", path)
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    engine = create_engine("sqlite:///:memory:")
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE gov_actions (id CHAR(32) PRIMARY KEY)"))
        connection.execute(text("CREATE TABLE user (id INTEGER PRIMARY KEY)"))
        connection.execute(text("INSERT INTO gov_actions (id) VALUES ('00000000000000000000000000000001')"))
        with Operations.context(MigrationContext.configure(connection)):
            migration.upgrade()
        assert connection.execute(text("SELECT tags FROM gov_actions")).scalar() == "[]"

        connection.execute(text("UPDATE gov_actions SET tags = '[\"Urgent\"]'"))
        with Operations.context(MigrationContext.configure(connection)):
            with pytest.raises(RuntimeError, match="tags"):
                migration.downgrade()
        assert "gov_action_evidence" in inspect(connection).get_table_names()
        connection.execute(text("UPDATE gov_actions SET tags = '[]'"))
        connection.execute(text("INSERT INTO user (id) VALUES (1)"))
        connection.execute(
            text(
                "INSERT INTO gov_action_evidence (id, created_at, action_id, storage_key, sha256, size_bytes, "
                "filename, content_type, uploaded_by_user_id) VALUES "
                "('00000000000000000000000000000002', '2026-09-24', '00000000000000000000000000000001', "
                "'action-evidence/test', 'abc', 1, 'proof.pdf', 'application/pdf', 1)"
            )
        )
        with Operations.context(MigrationContext.configure(connection)):
            with pytest.raises(RuntimeError, match="evidence metadata"):
                migration.downgrade()
        connection.execute(text("DELETE FROM gov_action_evidence"))
        with Operations.context(MigrationContext.configure(connection)):
            migration.downgrade()
        assert "tags" not in {column["name"] for column in inspect(connection).get_columns("gov_actions")}
    engine.dispose()
